from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title
    slide.placeholders[1].text = subtitle


def add_bullets_slide(prs, title, bullets):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    tf = slide.shapes.placeholders[1].text_frame
    tf.clear()
    for idx, bullet in enumerate(bullets):
        p = tf.add_paragraph() if idx > 0 else tf.paragraphs[0]
        p.text = bullet
        p.level = 0


def add_process_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Workflow Recommande"

    steps = [
        "1. Export JSON depuis Looker Studio (UI)",
        "2. Migrate avec migrate.py",
        "3. Ouvrir le .pbip dans Power BI Desktop",
        "4. Reconnecter BigQuery et appliquer les requetes",
        "5. Valider KPI/visuels puis publier",
    ]

    left = Inches(0.8)
    top = Inches(1.4)
    width = Inches(11.4)
    height = Inches(0.75)

    for i, step in enumerate(steps):
        y = top + Inches(i * 1.0)
        shape = slide.shapes.add_shape(1, left, y, width, height)  # rectangle
        fill = shape.fill
        fill.solid()
        if i % 2 == 0:
            fill.fore_color.rgb = RGBColor(10, 85, 140)
        else:
            fill.fore_color.rgb = RGBColor(18, 122, 91)
        shape.line.color.rgb = RGBColor(255, 255, 255)

        tf = shape.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = step
        p.font.size = Pt(20)
        p.font.bold = True
        p.font.color.rgb = RGBColor(255, 255, 255)


def build_presentation(output_path: Path):
    prs = Presentation()

    add_title_slide(
        prs,
        "Looker Studio -> Power BI",
        "Presentation de l'outil de migration\nEtat, workflow et limites",
    )

    add_bullets_slide(
        prs,
        "Ce que fait l'outil",
        [
            "Convertit des rapports Looker Studio (JSON) en projets Power BI (.pbip)",
            "Genere un package PBIP: Report + SemanticModel + dashboard de migration",
            "Supporte la migration batch de plusieurs rapports",
            "Prepare un flux reproductible pour les environnements GCP/BigQuery",
        ],
    )

    add_bullets_slide(
        prs,
        "Ce qui est deja en place",
        [
            "Generation de samples Looker Studio adaptes au schema BigQuery",
            "Sorties PBIP generees dans plusieurs dossiers de test",
            "Corrections successives du format PBIP/PBIR et des erreurs de parsing",
            "Connexion GCP validee via gcloud pour extraire schema et metadonnees BigQuery",
        ],
    )

    add_process_slide(prs)

    add_bullets_slide(
        prs,
        "Limites actuelles",
        [
            "API Looker Studio non exploitable ici pour lister/exporter tous les rapports",
            "Le chemin fiable reste l'export JSON via UI Looker Studio",
            "La reconnexion finale des sources BigQuery se fait dans Power BI Desktop",
        ],
    )

    add_bullets_slide(
        prs,
        "Plan d'execution recommande",
        [
            "1) Exporter tous les rapports cibles en JSON depuis Looker Studio",
            "2) Lancer migrate.py en batch dans le dossier cible",
            "3) Ouvrir les .pbip, corriger connexions/requetes, appliquer changements",
            "4) Valider KPI, filtres, relations, puis publier dans Power BI Service",
        ],
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))


if __name__ == "__main__":
    out = Path("artifacts") / "Presentation_Outil_Migration_Looker_to_PowerBI.pptx"
    build_presentation(out)
    print(f"PPTX_CREATED={out}")
